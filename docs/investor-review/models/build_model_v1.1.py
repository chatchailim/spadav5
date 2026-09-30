from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L

import sys
OUT=sys.argv[1] if len(sys.argv)>1 else "SPADA-DOC-FUND-003-v1.1-investor-model.xlsx"  # แล้วรัน recalc ของ LibreOffice เพื่อคำนวณค่า
FN="Arial"
f_norm=Font(name=FN,size=10); f_bold=Font(name=FN,size=10,bold=True)
f_title=Font(name=FN,size=14,bold=True); f_sub=Font(name=FN,size=10,italic=True,color="595959")
f_in=Font(name=FN,size=10,color="0000FF"); f_link=Font(name=FN,size=10,color="008000")
f_hdr=Font(name=FN,size=10,bold=True,color="FFFFFF")
fill_in=PatternFill("solid",fgColor="FFFF00"); fill_hdr=PatternFill("solid",fgColor="1F3864")
fill_sec=PatternFill("solid",fgColor="D9E1F2"); fill_warn=PatternFill("solid",fgColor="FCE4D6")
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
BAHT='#,##0;(#,##0);-'; PCT='0.0%;(0.0%);-'; PCT2='0.00%'
S1,S2,S3,S4,S5,S6="อ่านก่อน","สมมติฐาน","กระแสเงินสด","การแปลงสภาพ","ความไว","สรุป"
A=f"'{S2}'!"   # assumptions prefix

wb=Workbook()
ws0=wb.active; ws0.title=S1
ws=wb.create_sheet(S2); wc=wb.create_sheet(S3); wv=wb.create_sheet(S4); wsn=wb.create_sheet(S5); wsm=wb.create_sheet(S6)
wb.move_sheet(S6,offset=-4)  # summary after read-first

def put(w,ref,val,font=f_norm,fmt=None,fill=None,bold=False,wrap=False,align=None,border=False,unlock=False):
    c=w[ref]; c.value=val; c.font=Font(name=FN,size=font.size,bold=(bold or font.bold),italic=font.italic,color=font.color)
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if wrap or align: c.alignment=Alignment(wrap_text=wrap,vertical="top",horizontal=align)
    if border: c.border=bd
    if unlock: c.protection=Protection(locked=False)
    return c
def inp(w,ref,val,fmt=None): return put(w,ref,val,f_in,fmt,fill_in,border=True,unlock=True)
def sec(w,ref,text,span=5):
    put(w,ref,text,f_bold,fill=fill_sec)
    r=w[ref].row
    for i in range(2,2+span): w.cell(row=r,column=i).fill=fill_sec
def hdr(w,row,cols,labels):
    for col,lab in zip(cols,labels):
        c=w[f"{col}{row}"]; c.value=lab; c.font=f_hdr; c.fill=fill_hdr; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=bd

# ================= อ่านก่อน =================
ws0.sheet_view.showGridLines=False
put(ws0,"B1","แบบจำลองการเงิน Angel Round — SPADA / OneManOS",f_title)
put(ws0,"B2","NextWaver.Net Co., Ltd. · เอกสาร SPADA-DOC-FUND-003 ฉบับ 1.1 (ฉบับนักลงทุน) · วันที่ออก: [ใส่วันที่ส่ง]",f_sub)
lines=[
("วัตถุประสงค์","แสดงกระแสเงินสด 24 เดือน ผลของเงื่อนไขเงินกู้แปลงสภาพต่อสัดส่วนหุ้น และความไวต่อสมมติฐานสำคัญ เพื่อให้ผู้ลงทุนทดลองปรับตัวเลขเองได้"),
("วิธีใช้","แก้เฉพาะช่องที่เป็นตัวเลขสีน้ำเงินบนพื้นเหลืองในชีต 'สมมติฐาน' เท่านั้น ช่องสีดำเป็นสูตร ช่องสีเขียวเป็นการดึงค่าจากชีตอื่น ทุกชีตล็อกสูตรไว้ (ไม่มีรหัสผ่าน เพื่อป้องกันการพิมพ์ทับโดยไม่ตั้งใจ ไม่ใช่มาตรการความปลอดภัย)"),
("หน่วย","จำนวนเงินเป็นบาท อัตราส่วนเป็นร้อยละ เดือนที่ 1 = เดือนที่ได้รับงวดที่ 1"),
("ชีตในไฟล์","สรุป · สมมติฐาน · กระแสเงินสด (24 เดือน) · การแปลงสภาพ (ทั้งเพดานแบบ pre-money และ post-money) · ความไว (10 สถานการณ์)"),
("ข้อจำกัดสำคัญ","1) ตัวเลขทั้งหมดเป็นสมมติฐานเพื่อการวางแผน ไม่ใช่การรับประกันผลลัพธ์ ไม่ใช่คำเสนอขายหลักทรัพย์ ไม่ใช่คำแนะนำทางการเงินหรือกฎหมาย"),
("","2) ค่า pen-test และฮาร์ดแวร์ต้นแบบเป็นประมาณการ ยังไม่มีใบเสนอราคา ดูผลกระทบในชีต 'ความไว'"),
("","3) ยังไม่มีผู้ใช้จริงและยังไม่มีรายได้ แบบจำลองนี้ไม่มีประมาณการรายได้ เป็นแบบจำลองกระแสเงินสดและการแปลงสภาพเท่านั้น"),
("","4) สัญญาเงินกู้แปลงสภาพต้องร่างโดยทนายไทย ผลคำนวณในไฟล์นี้ใช้นิยามที่ระบุในชีตสมมติฐาน ซึ่งต้องตรงกับสัญญาฉบับจริง"),
("","5) ยังไม่รวมภาษีหัก ณ ที่จ่ายของดอกเบี้ย และค่าธรรมเนียมทางกฎหมายของฝ่ายนักลงทุน"),
("","6) การแปลงสภาพต้องผ่านมติผู้ถือหุ้นและการเพิ่มทุนตามกฎหมายบริษัทจำกัด"),
]
r=4
for k,v in lines:
    put(ws0,f"B{r}",k,f_bold,wrap=True); put(ws0,f"C{r}",v,f_norm,wrap=True); ws0.row_dimensions[r].height=44 if len(v)>110 else 30; r+=1
put(ws0,f"B{r+1}","สัญลักษณ์สี",f_bold)
put(ws0,f"C{r+1}","ตัวเลขสีน้ำเงินบนพื้นเหลือง = ช่องกรอก",f_in,fill=fill_in)
put(ws0,f"C{r+2}","สีดำ = สูตรคำนวณ",f_norm)
put(ws0,f"C{r+3}","สีเขียว = ดึงค่าจากชีตอื่น",f_link)
ws0.column_dimensions["A"].width=2; ws0.column_dimensions["B"].width=18; ws0.column_dimensions["C"].width=110

# ================= สมมติฐาน =================
ws.sheet_view.showGridLines=False
put(ws,"B1","สมมติฐาน (ช่องกรอกทั้งหมดอยู่ที่ชีตนี้)",f_title)
put(ws,"B2","แก้เฉพาะตัวเลขสีน้ำเงินบนพื้นเหลือง · ทุกจำนวนเงินเป็นบาท · ทุกอัตราส่วนเป็นร้อยละ",f_sub)
sec(ws,"B9","1 · การระดมทุน",3)
rows=[
(10,"เงินที่เสนอทั้งรอบ (บาท)",3000000,BAHT,"ยอดที่เสนอกับนักลงทุน"),
(11,"งวดที่ 1 (บาท)",2000000,BAHT,"โอนหลังลงนามสัญญา"),
]
for rr,lab,v,fmt,note in rows:
    put(ws,f"B{rr}",lab); inp(ws,f"C{rr}",v,fmt); put(ws,f"E{rr}",note,f_norm,wrap=True)
put(ws,"B12","งวดที่ 2 (บาท)"); put(ws,"C12","=C10-C11",fmt=BAHT); put(ws,"E12","ยอดรวมลบงวดที่ 1 · โอนเมื่อผ่านเงื่อนไขปลดงวดที่ 2",wrap=True)
put(ws,"B13","เดือนที่ได้รับงวดที่ 1"); inp(ws,"C13",1,'0'); put(ws,"E13","1 = เดือนแรกของแบบจำลอง")
put(ws,"B14","เดือนที่ได้รับงวดที่ 2"); inp(ws,"C14",7,'0'); put(ws,"E14","หลังผ่านเงื่อนไขและโอนภายใน 10 วันทำการ (ดูชีตสรุป)")
put(ws,"B15","เงินสดตั้งต้นในบริษัท (บาท)"); inp(ws,"C15",0,BAHT); put(ws,"E15","ใส่ยอดจริง ณ วันที่ส่งเอกสาร")
sec(ws,"B17","2 · ค่าใช้จ่ายประจำเดือน",3)
for rr,lab,v,note in [(18,"ค่าแรงทีม (รวมผู้ก่อตั้ง)",145000,"ต้องตรงกับงบของบริษัท"),(19,"ค่าโครงสร้างพื้นฐาน (คลาวด์)",3000,"โหนดไทย + สิงคโปร์ ตามประมาณการ"),(20,"ค่าดำเนินการอื่น",2000,"เครื่องมือ ค่าธรรมเนียม เบ็ดเตล็ด")]:
    put(ws,f"B{rr}",lab); inp(ws,f"C{rr}",v,BAHT); put(ws,f"E{rr}",note)
put(ws,"B21","รวมค่าใช้จ่ายต่อเดือน (บาท)",f_bold); put(ws,"C21","=SUM(C18:C20)",f_bold,BAHT)
put(ws,"B22","อัตราค่าใช้จ่ายเพิ่มขึ้นต่อเดือน"); inp(ws,"C22",0,PCT); put(ws,"E22","ใส่ 2% ถ้าคาดว่าจะจ้างเพิ่มระหว่างทาง")
sec(ws,"B24","3 · ค่าใช้จ่ายก้อนใหญ่ (จ่ายครั้งเดียว)",3)
hdr(ws,25,["B","C","D","E"],["รายการ","จำนวนเงิน (บาท)","เดือนที่จ่าย","หมายเหตุ"])
oneoffs=[(26,"กฎหมาย · โอนสิทธิ์ IP · จดทะเบียน",120000,1,"ดำเนินการก่อนลงนามสัญญาลงทุน"),
(27,"Pen-test รอบที่ 1",450000,2,"ประมาณการ ยังไม่มีใบเสนอราคา ช่วงราคาที่ประเมินไว้ 450,000–900,000 บาท ดูชีตความไว"),
(28,"Retest หลังแก้ไข",150000,6,"ต้องเสร็จก่อนปลดงวดที่ 2 (ใช้ยืนยันการแก้ช่องโหว่ Critical/High) ควรรวมในสัญญาผู้ทดสอบตั้งแต่แรก"),
(29,"ฮาร์ดแวร์ Setbox ชุดต้นแบบ",240000,9,"2–3 ชุดสำหรับชุมชนนำร่อง ราคายังไม่สรุป รอใบเสนอราคา")]
for rr,lab,amt,m,note in oneoffs:
    put(ws,f"B{rr}",lab); inp(ws,f"C{rr}",amt,BAHT); inp(ws,f"D{rr}",m,'0'); put(ws,f"E{rr}",note,f_norm,wrap=True)
put(ws,"B30","รวมค่าใช้จ่ายก้อนใหญ่ (บาท)",f_bold); put(ws,"C30","=SUM(C26:C29)",f_bold,BAHT)
sec(ws,"B32","4 · เงื่อนไขเงินกู้แปลงสภาพ",3)
put(ws,"B33","จำนวนหุ้นปัจจุบันของบริษัท"); inp(ws,"C33",100000,'#,##0'); put(ws,"E33","จำนวนหุ้นตัวอย่าง จะแทนด้วยจำนวนหุ้นจริงตาม บอจ.5 · สัดส่วนที่คำนวณไม่ขึ้นกับค่านี้",fill=fill_warn,wrap=True)
put(ws,"B34","เพดานมูลค่า valuation cap (บาท)"); inp(ws,"C34",15000000,BAHT)
put(ws,"B35","นิยามเพดาน (1 = pre-money, 2 = post-money)"); inp(ws,"C35",1,'0'); put(ws,"E35","1: สัดส่วน = ยอดแปลง ÷ (ยอดแปลง + เพดาน) · 2: สัดส่วน = ยอดแปลง ÷ เพดาน (ต้องตรงกับสัญญา)",wrap=True)
put(ws,"B36","ส่วนลดจากราคารอบถัดไป"); inp(ws,"C36",0.2,PCT)
put(ws,"B37","ดอกเบี้ยต่อปี (คงที่ สะสม ไม่จ่ายเป็นเงินสด)"); inp(ws,"C37",0.03,PCT); put(ws,"E37","คิดจากวันที่เบิกจริงของแต่ละงวด · ภาษีหัก ณ ที่จ่ายยังไม่รวม")
put(ws,"B38","เดือนที่แปลงสภาพ/ครบอายุสัญญา (นับจากเดือนที่ 1)"); inp(ws,"C38",24,'0'); put(ws,"E38","24 = อายุสัญญา 24 เดือน")
put(ws,"B39","มูลค่าก่อนเงินของรอบถัดไป (บาท)"); inp(ws,"C39",40000000,BAHT); put(ws,"E39","สมมติฐานเพื่อทดสอบว่าเพดานหรือส่วนลดจะมีผลบังคับ ไม่ใช่การคาดการณ์มูลค่า",wrap=True)
put(ws,"B41","ข้อจำกัดของแบบจำลองนี้",f_bold)
for i,t in enumerate(["• ตัวเลขทั้งหมดเป็นสมมติฐานเพื่อการวางแผน ไม่ใช่ใบเสนอราคาหรือคำแนะนำทางการเงิน","• สัญญาเงินกู้แปลงสภาพในไทยต้องให้ทนายไทยร่าง","• การแปลงสภาพต้องผ่านมติผู้ถือหุ้นและการเพิ่มทุนตามกฎหมายบริษัทจำกัด","• ยังไม่รวมภาษีหัก ณ ที่จ่ายของดอกเบี้ย และค่าธรรมเนียมทางกฎหมายของฝ่ายนักลงทุน"]):
    put(ws,f"B{42+i}",t)
for col,w in zip("ABCDE",[2,52,20,14,80]): ws.column_dimensions[col].width=w

# ================= กระแสเงินสด =================
wc.sheet_view.showGridLines=False
put(wc,"B1","กระแสเงินสด 24 เดือน",f_title); put(wc,"B2","ทุกช่องเป็นสูตร · แก้ค่าได้ที่ชีตสมมติฐานเท่านั้น",f_sub)
hdr(wc,4,list("BCDEFGHIJKL"),["เดือนที่","เงินสดต้นงวด","เงินลงทุนเข้า","ค่าแรง","คลาวด์","อื่นๆ","ก้อนใหญ่","รวมจ่าย","เงินสดปลายงวด","สถานะ","ตัวช่วย (ติดลบ=1)"])
for m in range(1,25):
    r=4+m
    put(wc,f"B{r}",m,fmt='0')
    put(wc,f"C{r}",f"={A}$C$15" if m==1 else f"=J{r-1}",f_link if m==1 else f_norm,BAHT)
    put(wc,f"D{r}",f"=IF($B{r}={A}$C$13,{A}$C$11,0)+IF($B{r}={A}$C$14,{A}$C$12,0)",f_link,BAHT)
    put(wc,f"E{r}",f"={A}$C$18*(1+{A}$C$22)^($B{r}-1)",f_link,BAHT)
    put(wc,f"F{r}",f"={A}$C$19*(1+{A}$C$22)^($B{r}-1)",f_link,BAHT)
    put(wc,f"G{r}",f"={A}$C$20*(1+{A}$C$22)^($B{r}-1)",f_link,BAHT)
    put(wc,f"H{r}",f"=SUMIF({A}$D$26:$D$29,$B{r},{A}$C$26:$C$29)",f_link,BAHT)
    put(wc,f"I{r}",f"=SUM(E{r}:H{r})",fmt=BAHT)
    put(wc,f"J{r}",f"=C{r}+D{r}-I{r}",f_bold,BAHT)
    put(wc,f"K{r}",f'=IF(J{r}<0,"เงินหมด",IF(J{r}<{A}$C$21*2,"เหลือน้อย","ปกติ"))')
    put(wc,f"L{r}",f"=IF(J{r}<0,1,0)",fmt='0')
    for col in "BCDEFGHIJKL": wc[f"{col}{r}"].border=bd
sec(wc,"B31","สรุปกระแสเงินสด",4)
put(wc,"B32","เดือนแรกที่เงินสดติดลบ (0 = ไม่ติดลบใน 24 เดือน)"); put(wc,"F32","=IFERROR(MATCH(1,L5:L28,0),0)",f_bold,'0')
put(wc,"B33","จำนวนเดือนที่เงินยังไม่หมด (runway)"); put(wc,"F33","=IF(F32=0,24,F32-1)",f_bold,'0')
put(wc,"B34","เงินสดต่ำสุดใน 18 เดือนแรก (บาท)"); put(wc,"F34","=MIN(J5:J22)",f_bold,BAHT)
put(wc,"B35","เงินที่ต้องเพิ่มเพื่อให้ไม่ติดลบตลอด 18 เดือน (บาท)"); put(wc,"F35","=MAX(0,-F34)",f_bold,BAHT)
put(wc,"B36","เงินสดสิ้นเดือนที่ 6 (ก่อนงวดที่ 2) (บาท)"); put(wc,"F36","=J10",f_bold,BAHT)
for col,w in zip("ABCDEFGHIJKL",[2,16,16,16,12,10,10,14,14,16,12,16]): wc.column_dimensions[col].width=w
wc.freeze_panes="C5"

# ================= การแปลงสภาพ =================
wv.sheet_view.showGridLines=False
put(wv,"B1","การแปลงสภาพเป็นหุ้น และสัดส่วนที่นักลงทุนได้",f_title); put(wv,"B2","แก้เพดาน ส่วนลด ดอกเบี้ย และนิยามเพดานได้ที่ชีตสมมติฐาน แถว 34–39",f_sub)
sec(wv,"B4","ก. ยอดที่จะแปลงสภาพ (ดอกเบี้ยคิดจากวันที่เบิกจริงของแต่ละงวด)",4)
hdr(wv,5,list("BCDE"),["รายการ","งวดที่ 1","งวดที่ 2","รวม"])
put(wv,"B6","เงินต้น (บาท)"); put(wv,"C6",f"={A}C11",f_link,BAHT); put(wv,"D6",f"={A}C12",f_link,BAHT); put(wv,"E6","=C6+D6",fmt=BAHT)
put(wv,"B7","เดือนที่เบิก"); put(wv,"C7",f"={A}C13",f_link,'0'); put(wv,"D7",f"={A}C14",f_link,'0')
put(wv,"B8","จำนวนเดือนที่ดอกเบี้ยสะสม"); put(wv,"C8",f"=MAX(0,{A}$C$38-(C7-1))",f_link,'0'); put(wv,"D8",f"=MAX(0,{A}$C$38-(D7-1))",f_link,'0')
put(wv,"B9","ดอกเบี้ยสะสม (บาท)"); put(wv,"C9",f"=C6*{A}$C$37*C8/12",f_link,BAHT); put(wv,"D9",f"=D6*{A}$C$37*D8/12",f_link,BAHT); put(wv,"E9","=C9+D9",fmt=BAHT)
put(wv,"B10","ยอดที่นำไปแปลงเป็นหุ้น (บาท)",f_bold); put(wv,"C10","=C6+C9",f_bold,BAHT); put(wv,"D10","=D6+D9",f_bold,BAHT); put(wv,"E10","=E6+E9",f_bold,BAHT)
sec(wv,"B12","ข. ราคาต่อหุ้น: เพดานกับส่วนลด อันไหนใช้จริง (คิดแบบ pre-money)",4)
put(wv,"B13","ราคาต่อหุ้นจากเพดานมูลค่า (บาท)"); put(wv,"C13",f"={A}C34/{A}C33",f_link,'#,##0.00')
put(wv,"B14","ราคาต่อหุ้นรอบถัดไป ก่อนส่วนลด (บาท)"); put(wv,"C14",f"={A}C39/{A}C33",f_link,'#,##0.00')
put(wv,"B15","ราคาต่อหุ้นรอบถัดไป หลังส่วนลด (บาท)"); put(wv,"C15",f"=C14*(1-{A}C36)",f_link,'#,##0.00')
put(wv,"B16","ราคาที่ใช้จริง (ต่ำกว่าชนะ) (บาท)"); put(wv,"C16","=MIN(C13,C15)",fmt='#,##0.00')
put(wv,"B17","ข้อไหนมีผลบังคับ"); put(wv,"C17",'=IF(C13<=C15,"เพดานมูลค่า","ส่วนลด")',f_bold)
put(wv,"B18","ส่วนลดจะมีผลเมื่อมูลค่าก่อนเงินของรอบถัดไปต่ำกว่า (บาท)"); put(wv,"C18",f"={A}C34/(1-{A}C36)",f_link,BAHT)
sec(wv,"B20","ค. ผลต่อสัดส่วนผู้ถือหุ้น (ตามนิยามเพดานที่เลือกในสมมติฐาน)",4)
put(wv,"B21","สัดส่วนนักลงทุน ถ้าใช้เพดาน"); put(wv,"C21",f"=IF({A}C35=1,E10/(E10+{A}C34),E10/{A}C34)",f_link,PCT2)
put(wv,"B22","สัดส่วนนักลงทุน ถ้าใช้ส่วนลด"); put(wv,"C22",f"=E10/(E10+{A}C39*(1-{A}C36))",f_link,PCT2)
put(wv,"B23","สัดส่วนที่ใช้จริง (ค่าที่ให้นักลงทุนมากกว่า)",f_bold); put(wv,"C23","=MAX(C21,C22)",f_bold,PCT2)
put(wv,"B24","จำนวนหุ้นที่นักลงทุนจะได้"); put(wv,"C24",f"=C23/(1-C23)*{A}C33",f_link,'#,##0')
put(wv,"B25","จำนวนหุ้นเดิม"); put(wv,"C25",f"={A}C33",f_link,'#,##0')
put(wv,"B26","จำนวนหุ้นรวมหลังแปลงสภาพ"); put(wv,"C26","=C24+C25",fmt='#,##0')
put(wv,"B27","สัดส่วนที่ผู้ถือหุ้นเดิมเหลือ",f_bold); put(wv,"C27","=1-C23",f_bold,PCT2)
put(wv,"B28","ตรวจสอบ: หุ้นที่นักลงทุนได้ ÷ หุ้นรวม (ต้องเท่ากับสัดส่วนที่ใช้จริง)"); put(wv,"C28","=C24/C26",fmt=PCT2)
sec(wv,"B30","ง. ความไวของเพดานมูลค่า (ยอดแปลงตามข้อ ก)",4)
hdr(wv,31,list("BCDE"),["เพดานมูลค่า (บาท)","สัดส่วนนักลงทุน ถ้าเพดานเป็น pre-money","สัดส่วนนักลงทุน ถ้าเพดานเป็น post-money","เพดาน post-money ที่ให้สัดส่วนเท่ากับ pre-money นี้ (บาท)"])
wv.row_dimensions[31].height=54
for i,cap in enumerate([8e6,10e6,12e6,15e6,18e6,20e6,25e6,30e6]):
    r=32+i; inp(wv,f"B{r}",cap,BAHT)
    put(wv,f"C{r}",f"=$E$10/($E$10+B{r})",fmt=PCT2,border=True); put(wv,f"D{r}",f"=IF(B{r}>0,$E$10/B{r},0)",fmt=PCT2,border=True); put(wv,f"E{r}",f"=$E$10+B{r}",fmt=BAHT,border=True)
sec(wv,"B41","จ. ความไวของดอกเบี้ย (ที่เพดานและนิยามที่เลือก)",4)
hdr(wv,42,list("BCD"),["ดอกเบี้ยต่อปี","ยอดแปลง (บาท)","สัดส่วนนักลงทุน"])
for i,rate in enumerate([0.0,0.03,0.05,0.08]):
    r=43+i; inp(wv,f"B{r}",rate,PCT)
    put(wv,f"C{r}",f"=$E$6+B{r}*($C$6*$C$8+$D$6*$D$8)/12",fmt=BAHT,border=True)
    put(wv,f"D{r}",f"=IF({A}$C$35=1,C{r}/(C{r}+{A}$C$34),C{r}/{A}$C$34)",f_link,PCT2,border=True)
sec(wv,"B48","ฉ. ภาระเงินสดถ้าตราสารไม่แปลงสภาพและไม่ต่ออายุ (ณ เดือนที่แปลงสภาพ/ครบอายุ)",4)
put(wv,"B49","เบิกครบทั้งสองงวด: เงินต้น + ดอกเบี้ย (บาท)"); put(wv,"C49","=E10",fmt=BAHT)
put(wv,"B50","ได้เฉพาะงวดที่ 1: เงินต้น + ดอกเบี้ย (บาท)"); put(wv,"C50","=C10",fmt=BAHT)
for col,w in zip("ABCDE",[2,60,26,26,32]): wv.column_dimensions[col].width=w

# ================= ความไว =================
wsn.sheet_view.showGridLines=False
put(wsn,"B1","ความไวของกระแสเงินสดต่อสมมติฐานสำคัญ (10 สถานการณ์)",f_title)
put(wsn,"B2","แถวแรกดึงค่าจากชีตสมมติฐาน แถวอื่นเปลี่ยนเฉพาะช่องสีน้ำเงิน ที่เหลือเท่ากับแถวฐาน · เงินสดรายเดือนคำนวณด้วยสูตรแบบปิดรูป (ผลเท่ากับชีตกระแสเงินสดในแถวฐาน)",f_sub)
hdr(wsn,5,list("BCDEFGH")+list("IJKLMNO"),["สถานการณ์","งวดที่ 1 (บาท)","งวดที่ 2 (บาท)","เดือนที่ได้งวดที่ 2","Pen-test (บาท)","ฮาร์ดแวร์ (บาท)","ค่าใช้จ่ายเพิ่มต่อเดือน","เดือนแรกที่ติดลบ (0=ไม่ติดลบ)","runway (เดือน)","เงินสดสิ้นเดือน 6","เงินสดสิ้นเดือน 7","เงินสดสิ้นเดือน 8","เงินสดต่ำสุด 18 เดือน","เงินที่ต้องเพิ่มเพื่อครบ 18 เดือน"])
wsn.row_dimensions[5].height=54
scen=[("ฐาน (ตามสมมติฐาน)",{}),("ได้เฉพาะงวดที่ 1 (ไม่ปลดงวดที่ 2)",{"D":0}),("งวดที่ 2 ล่าช้า 1 เดือน",{"E":"=$E$6+1"}),("งวดที่ 2 ล่าช้า 2 เดือน",{"E":"=$E$6+2"}),("Pen-test 700,000",{"F":700000}),("Pen-test 830,000 (เงินสดเดือน 6 ≈ 0)",{"F":830000}),("Pen-test 900,000",{"F":900000}),("ค่าใช้จ่ายเพิ่ม 2% ต่อเดือน",{"H":0.02}),("ฮาร์ดแวร์ 400,000",{"G":400000}),("Pen-test 900,000 + งวดที่ 2 ล่าช้า 1 เดือน",{"F":900000,"E":"=$E$6+1"})]
base_map={"C":f"={A}C11","D":f"={A}C12","E":f"={A}C14","F":f"={A}C27","G":f"={A}C29","H":f"={A}C22"}
fmts={"C":BAHT,"D":BAHT,"E":'0',"F":BAHT,"G":BAHT,"H":PCT}
GR0,FR0=20,34   # cash grid first row, flag grid first row
put(wsn,"B18","เงินสดปลายงวดรายเดือน (บาท)",f_bold); put(wsn,"B32","ตัวช่วย: ธงเงินสดติดลบ (1 = ติดลบ)",f_bold)
put(wsn,"B19","สถานการณ์",f_bold); put(wsn,"B33","สถานการณ์",f_bold)
for mi in range(24):
    col=L(3+mi)
    put(wsn,f"{col}19",mi+1,f_bold,'0',fill_sec,align="center"); put(wsn,f"{col}33",mi+1,f_bold,'0',fill_sec,align="center")
for i,(name,ov) in enumerate(scen):
    p=6+i; put(wsn,f"B{p}",name,f_bold if i==0 else f_norm,border=True,wrap=True)
    for col in "CDEFGH":
        if col in ov:
            v=ov[col]
            if isinstance(v,str): put(wsn,f"{col}{p}",v,f_norm,fmts[col],border=True)
            else: inp(wsn,f"{col}{p}",v,fmts[col])
        else:
            if i==0: put(wsn,f"{col}{p}",base_map[col],f_link,fmts[col],border=True)
            else: put(wsn,f"{col}{p}",f"=${col}$6",f_norm,fmts[col],border=True)
    gr=GR0+i; fr=FR0+i
    put(wsn,f"B{gr}",f"=B{p}",wrap=True); put(wsn,f"B{fr}",f"=B{p}",wrap=True)
    for mi in range(24):
        col=L(3+mi); M=f"{col}$19"
        cash=(f"={A}$C$15+$C{p}*({M}>={A}$C$13)+$D{p}*({M}>=$E{p})"
              f"-IF($H{p}=0,{A}$C$21*{M},{A}$C$21*((1+$H{p})^{M}-1)/$H{p})"
              f"-({A}$C$26*({M}>={A}$D$26)+$F{p}*({M}>={A}$D$27)+{A}$C$28*({M}>={A}$D$28)+$G{p}*({M}>={A}$D$29))")
        put(wsn,f"{col}{gr}",cash,f_norm,BAHT)
        put(wsn,f"{col}{fr}",f"=IF({col}{gr}<0,1,0)",f_norm,'0')
    put(wsn,f"I{p}",f"=IFERROR(MATCH(1,C{fr}:Z{fr},0),0)",f_bold,'0',border=True)
    put(wsn,f"J{p}",f"=IF(I{p}=0,24,I{p}-1)",f_bold,'0',border=True)
    put(wsn,f"K{p}",f"=INDEX(C{gr}:Z{gr},6)",fmt=BAHT,border=True)
    put(wsn,f"L{p}",f"=INDEX(C{gr}:Z{gr},7)",fmt=BAHT,border=True)
    put(wsn,f"M{p}",f"=INDEX(C{gr}:Z{gr},8)",fmt=BAHT,border=True)
    put(wsn,f"N{p}",f"=MIN(C{gr}:T{gr})",fmt=BAHT,border=True)
    put(wsn,f"O{p}",f"=MAX(0,-N{p})",fmt=BAHT,border=True)
    wsn.row_dimensions[p].height=28
put(wsn,"B17","หมายเหตุ: ทุกสถานการณ์ใช้ค่าอื่นเท่ากับแถวฐาน (ค่าใช้จ่ายรายเดือน ก้อนใหญ่อื่น เดือนที่จ่าย) · เดือนแรกที่ติดลบ = 0 หมายถึงไม่ติดลบใน 24 เดือน",f_sub)
wsn.column_dimensions["A"].width=2; wsn.column_dimensions["B"].width=40
for mi in range(24): wsn.column_dimensions[L(3+mi)].width=14
wsn.freeze_panes="C6"

# ================= สรุป =================
wsm.sheet_view.showGridLines=False
put(wsm,"B1","สรุปสำหรับการเจรจา",f_title); put(wsm,"B2","หน้านี้ดึงตัวเลขจากชีตอื่นอัตโนมัติ · ใช้เปิดคู่กับการหารือ",f_sub)
sec(wsm,"B4","ตัวเลขหลัก",3)
items=[("เงินที่เสนอทั้งรอบ (บาท)",f"={A}C10",BAHT),("งวดที่ 1 (บาท)",f"={A}C11",BAHT),("งวดที่ 2 (บาท)",f"={A}C12",BAHT),
("ค่าใช้จ่ายประจำต่อเดือน (บาท)",f"={A}C21",BAHT),("ค่าใช้จ่ายก้อนใหญ่รวม (บาท)",f"={A}C30",BAHT),
("อยู่ได้กี่เดือน ถ้าปลดงวดที่ 2 (เดือน)",f"='{S3}'!F33",'0'),("อยู่ได้กี่เดือน ถ้าได้เฉพาะงวดที่ 1 (เดือน)",f"='{S5}'!J7",'0'),
("เดือนแรกที่เงินติดลบ (ปลดงวดที่ 2)",f"='{S3}'!F32",'0'),("เงินสดสิ้นเดือน 6 ก่อนงวดที่ 2 (บาท)",f"='{S3}'!F36",BAHT),
("เงินที่ต้องเพิ่มเพื่อไม่ติดลบตลอด 18 เดือน (บาท)",f"='{S3}'!F35",BAHT),
("ยอดที่นำไปแปลงสภาพ (เงินต้น + ดอกเบี้ย) (บาท)",f"='{S4}'!E10",BAHT),
("สัดส่วนนักลงทุนหลังแปลงสภาพ (ตามนิยามเพดานที่เลือก)",f"='{S4}'!C23",PCT),
("สัดส่วนที่ผู้ถือหุ้นเดิมเหลือ",f"='{S4}'!C27",PCT),
("ข้อไหนมีผลบังคับ (เพดานหรือส่วนลด) ที่มูลค่ารอบถัดไปตามสมมติฐาน",f"='{S4}'!C17",None)]
for i,(lab,f,fmt) in enumerate(items):
    r=5+i; put(wsm,f"B{r}",lab,border=True); put(wsm,f"C{r}",f,f_link,fmt,border=True)
    if lab.startswith("ข้อไหน"): wsm[f"C{r}"].alignment=Alignment(horizontal="right")
r0=5+len(items)+1
sec(wsm,f"B{r0}","เงื่อนไขปลดงวดที่ 2 (เขียนลงสัญญา)",3)
for i,t in enumerate(["• รายงานการทดสอบเจาะระบบจากผู้ตรวจอิสระที่ระบุชื่อในสัญญาออกแล้ว","• ช่องโหว่ระดับ Critical และ High ได้รับการแก้ไข พร้อมหนังสือยืนยันผลการทดสอบซ้ำจากผู้ตรวจ","• มีผู้ใช้จริงอย่างน้อย 10 ราย บนระบบที่เปิดใช้งาน (นิยาม \"ผู้ใช้จริง\" ระบุในสัญญา)","• ส่งรายงานความคืบหน้าให้นักลงทุนครบทุกไตรมาส","• นักลงทุนโอนงวดที่ 2 ภายใน 10 วันทำการหลังผู้ตรวจอิสระออกหนังสือยืนยันเงื่อนไข"]):
    put(wsm,f"B{r0+1+i}",t)
r1=r0+7
sec(wsm,f"B{r1}","วัดผลทั้งรอบ",3)
put(wsm,f"B{r1+1}","• มีธุรกิจที่จ่ายเงินจริงอย่างน้อย 5 ราย ภายในเดือนที่ 12")
put(wsm,f"B{r1+2}","• ถ้าไม่ถึง แปลว่าโมเดลรายได้ผิด ไม่ใช่เงินไม่พอ — และบริษัทจะเปลี่ยนแผน ไม่ใช่ขอเงินเพิ่ม")
put(wsm,f"B{r1+4}","แบบจำลองนี้เป็นสมมติฐานเพื่อการวางแผน ไม่ใช่คำแนะนำทางการเงินหรือกฎหมาย ไม่ใช่คำเสนอขายหลักทรัพย์",f_sub)
wsm.column_dimensions["A"].width=2; wsm.column_dimensions["B"].width=70; wsm.column_dimensions["C"].width=24

for w in wb.worksheets:
    w.protection.sheet=True
wb.save(OUT); print("saved",OUT)
